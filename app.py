import sqlite3, re, os, sys
from pathlib import Path
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook, Workbook

APP_DIR = Path(__file__).resolve().parent

def app_data_dir():
    # IMPORTANT: PyInstaller --onefile extracts the EXE to a temporary folder.
    # Never store the SQLite database beside __file__ when running frozen.
    if getattr(sys, 'frozen', False):
        root = Path(os.environ.get('LOCALAPPDATA') or Path.home() / 'AppData' / 'Local')
        return root / 'Noor-ul-Hifz'
    return APP_DIR / 'data'

DATA_DIR = app_data_dir()
DB_PATH = DATA_DIR / 'student_logs.db'
RANGES = [
    ('absent', 4, 7),
    ('takheer', 11, 22),
    ('sabaq', 26, 39),
    ('sabaqi', 43, 54),
    ('manzil', 58, 69),
]
LABELS = {
    'absent': 'غیر حاضری',
    'takheer': 'تاخیر',
    'sabaq': 'سبق کا ناغہ',
    'sabaqi': 'سبقی کا ناغہ',
    'manzil': 'منزل کا ناغہ',
}

# Visual language is intentionally aligned with the Noor-ul-Hifz website:
# warm white surfaces, deep ink typography, restrained blue accents and generous space.
LIGHT = {
    'bg': '#F6F4EF', 'surface': '#FFFDF9', 'surface2': '#F0EEE8',
    'border': '#E4E0D8', 'text': '#17232D', 'muted': '#737B82',
    'accent': '#4776D8', 'accent_dark': '#345EB6', 'accent_soft': '#EAF0FC',
    'success': '#2F8061', 'danger': '#C95555', 'warning': '#B98232',
    'sidebar': '#FBFAF6', 'sidebar_hover': '#F0EEE8', 'table_head': '#F3F1EC',
    'chip': '#ECE9E2', 'ink_soft': '#D9D5CC'
}
DARK = {
    'bg': '#11171D', 'surface': '#171E25', 'surface2': '#1D252D',
    'border': '#2A343D', 'text': '#F5F2EA', 'muted': '#9AA3AA',
    'accent': '#7198E8', 'accent_dark': '#5D82CF', 'accent_soft': '#202D43',
    'success': '#55A986', 'danger': '#E27878', 'warning': '#D5A653',
    'sidebar': '#141A20', 'sidebar_hover': '#1D252D', 'table_head': '#1B232B',
    'chip': '#20282F', 'ink_soft': '#34404A'
}


def norm(v):
    if v is None:
        return ''
    s = str(v).strip()
    return s[:-2] if re.fullmatch(r'-?\d+\.0', s) else s


def conn():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def init_db():
    with conn() as c:
        c.execute('''CREATE TABLE IF NOT EXISTS daily_logs(
            log_date TEXT, student_id TEXT, student_name TEXT DEFAULT '',
            absent INTEGER DEFAULT 0, takheer INTEGER DEFAULT 0,
            sabaq INTEGER DEFAULT 0, sabaqi INTEGER DEFAULT 0,
            manzil INTEGER DEFAULT 0, created_at TEXT,
            PRIMARY KEY(log_date, student_id))''')
        c.execute('''CREATE TABLE IF NOT EXISTS students(
            student_id TEXT PRIMARY KEY, student_name TEXT DEFAULT '', updated_at TEXT)''')


def migrate_legacy_db():
    # If a development/portable database exists, keep it when moving to the
    # persistent Windows AppData location. Never overwrite a newer database.
    if not getattr(sys, 'frozen', False):
        return
    legacy = APP_DIR / 'data' / 'student_logs.db'
    if legacy.exists() and not DB_PATH.exists():
        try:
            DB_PATH.parent.mkdir(parents=True, exist_ok=True)
            DB_PATH.write_bytes(legacy.read_bytes())
        except Exception:
            pass



def read_report(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if 'یومیہ رپورٹ' not in wb.sheetnames:
            raise ValueError("Sheet 'یومیہ رپورٹ' was not found.")
        ws = wb['یومیہ رپورٹ']
        raw = ws['G1'].value
        if isinstance(raw, (datetime, date)):
            source_date = raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat()
        else:
            source_date = norm(raw) or date.today().isoformat()
        found = {}
        counts = {k: 0 for k in LABELS}
        for condition, start, end in RANGES:
            for r in range(start, end + 1):
                for id_col, name_col in [('B', 'C'), ('F', 'G')]:
                    sid = norm(ws[f'{id_col}{r}'].value)
                    name = norm(ws[f'{name_col}{r}'].value)
                    if not sid:
                        continue
                    if sid not in found:
                        found[sid] = {'student_id': sid, 'student_name': name, **{k: 0 for k in LABELS}}
                    elif name and not found[sid]['student_name']:
                        found[sid]['student_name'] = name
                    if not found[sid][condition]:
                        found[sid][condition] = 1
                        counts[condition] += 1
        return source_date, list(found.values()), counts
    finally:
        wb.close()


def save_records(log_date, recs):
    now = datetime.now().isoformat(timespec='seconds')
    with conn() as c:
        for r in recs:
            c.execute('''INSERT INTO daily_logs(
                log_date,student_id,student_name,absent,takheer,sabaq,sabaqi,manzil,created_at)
                VALUES(?,?,?,?,?,?,?,?,?)
                ON CONFLICT(log_date,student_id) DO UPDATE SET
                student_name=CASE WHEN excluded.student_name<>'' THEN excluded.student_name ELSE daily_logs.student_name END,
                absent=excluded.absent,
                takheer=excluded.takheer,
                sabaq=excluded.sabaq,
                sabaqi=excluded.sabaqi,
                manzil=excluded.manzil,
                created_at=excluded.created_at''',
                (log_date, r['student_id'], r['student_name'], r['absent'], r['takheer'], r['sabaq'], r['sabaqi'], r['manzil'], now))
            c.execute('''INSERT INTO students(student_id,student_name,updated_at) VALUES(?,?,?)
                ON CONFLICT(student_id) DO UPDATE SET
                student_name=CASE WHEN excluded.student_name<>'' THEN excluded.student_name ELSE students.student_name END,
                updated_at=excluded.updated_at''',
                (r['student_id'], r['student_name'], now))


def export_student(sid, name, path):
    with conn() as c:
        rows = c.execute('SELECT * FROM daily_logs WHERE student_id=? ORDER BY log_date', (sid,)).fetchall()
    wb = Workbook(); ws = wb.active; ws.title = 'Student Report'
    ws.append(['ID','Name','Date','غیر حاضری','تاخیر','سبق کا ناغہ','سبقی کا ناغہ','منزل کا ناغہ'])
    for r in rows:
        ws.append([r['student_id'], r['student_name'] or name, r['log_date'], r['absent'], r['takheer'], r['sabaq'], r['sabaqi'], r['manzil']])
    ws.append([])
    ws.append(['Totals','','',sum(r['absent'] for r in rows),sum(r['takheer'] for r in rows),sum(r['sabaq'] for r in rows),sum(r['sabaqi'] for r in rows),sum(r['manzil'] for r in rows)])
    ws.freeze_panes = 'A2'; wb.save(path)


def totals():
    with conn() as c:
        students = c.execute('SELECT COUNT(*) n FROM students').fetchone()['n']
        logs = c.execute('SELECT COUNT(*) n FROM daily_logs').fetchone()['n']
        days = c.execute('SELECT COUNT(DISTINCT log_date) n FROM daily_logs').fetchone()['n']
        today = c.execute('''SELECT
            COALESCE(SUM(absent),0) absent, COALESCE(SUM(takheer),0) takheer,
            COALESCE(SUM(sabaq),0) sabaq, COALESCE(SUM(sabaqi),0) sabaqi,
            COALESCE(SUM(manzil),0) manzil FROM daily_logs WHERE log_date=?''',
            (date.today().isoformat(),)).fetchone()
        return students, logs, days, today


def recent_days(limit=7):
    with conn() as c:
        return c.execute('''SELECT log_date, COUNT(*) students,
            COALESCE(SUM(absent),0) absent, COALESCE(SUM(takheer),0) takheer,
            COALESCE(SUM(sabaq),0) sabaq, COALESCE(SUM(sabaqi),0) sabaqi,
            COALESCE(SUM(manzil),0) manzil
            FROM daily_logs GROUP BY log_date ORDER BY log_date DESC LIMIT ?''', (limit,)).fetchall()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Noor-ul-Hifz — نور الحفظ')
        self.geometry('1320x800'); self.minsize(1120,700)
        self.theme = LIGHT
        self.current = None
        self.status_message = 'Offline · Local database'
        self.setup_style(); migrate_legacy_db(); init_db(); self.build_shell(); self.show_page('dashboard')

    def setup_style(self):
        st = ttk.Style(self)
        try: st.theme_use('clam')
        except tk.TclError: pass
        st.configure('TButton', font=('Segoe UI', 10), padding=(13, 8), relief='flat', borderwidth=0)
        st.configure('Accent.TButton', font=('Segoe UI', 10, 'bold'), padding=(15, 9), relief='flat', borderwidth=0)
        st.configure('TEntry', font=('Segoe UI', 10), padding=8, relief='flat')
        st.configure('Treeview', font=('Segoe UI', 10), rowheight=44, borderwidth=0)
        st.configure('Treeview.Heading', font=('Segoe UI', 9, 'bold'), padding=10, relief='flat')

    def apply_tree_style(self):
        st = ttk.Style(self)
        st.configure('Treeview', background=self.theme['surface'], fieldbackground=self.theme['surface'], foreground=self.theme['text'], rowheight=44, borderwidth=0)
        st.map('Treeview', background=[('selected', self.theme['accent'])], foreground=[('selected', '#FFFFFF')])
        st.configure('Treeview.Heading', background=self.theme['table_head'], foreground=self.theme['muted'], font=('Segoe UI', 9, 'bold'))
        st.configure('TButton', background=self.theme['surface'], foreground=self.theme['text'])
        st.map('TButton', background=[('active', self.theme['sidebar_hover'])], foreground=[('active', self.theme['text'])])
        st.configure('Accent.TButton', background=self.theme['accent'], foreground='#FFFFFF')
        st.map('Accent.TButton', background=[('active', self.theme['accent_dark'])])
        st.configure('TEntry', fieldbackground=self.theme['surface'], foreground=self.theme['text'], bordercolor=self.theme['border'])

    def build_shell(self):
        self.configure(bg=self.theme['bg'])
        self.sidebar = tk.Frame(self, bg=self.theme['sidebar'], width=238, highlightthickness=1, highlightbackground=self.theme['border'])
        self.sidebar.pack(side='left', fill='y'); self.sidebar.pack_propagate(False)
        self.content = tk.Frame(self, bg=self.theme['bg']); self.content.pack(side='left', fill='both', expand=True)

        brand = tk.Frame(self.sidebar, bg=self.theme['sidebar']); brand.pack(fill='x', padx=20, pady=(24,30))
        logo = tk.Frame(brand, bg=self.theme['accent'], width=42, height=42); logo.pack(side='left'); logo.pack_propagate(False)
        tk.Label(logo, text='ن', bg=self.theme['accent'], fg='white', font=('Segoe UI', 18, 'bold')).pack(expand=True)
        bt = tk.Frame(brand, bg=self.theme['sidebar']); bt.pack(side='left', padx=11)
        tk.Label(bt, text='Noor-ul-Hifz', bg=self.theme['sidebar'], fg=self.theme['text'], font=('Segoe UI', 12, 'bold')).pack(anchor='w')
        tk.Label(bt, text='نور الحفظ', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI', 9)).pack(anchor='w', pady=(1,0))

        tk.Label(self.sidebar, text='WORKSPACE', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI', 8, 'bold')).pack(anchor='w', padx=22, pady=(0,8))
        self.nav_buttons = {}
        nav = [
            ('dashboard','Dashboard','▦'),
            ('import','Import Daily Report','↥'),
            ('manual','Manual Entry','＋'),
            ('report','Collect Report','◎'),
            ('all','All Students','▤'),
        ]
        for key, title, icon in nav:
            b = tk.Button(self.sidebar, text=f'  {icon}   {title}', command=lambda k=key:self.show_page(k), anchor='w', relief='flat', bd=0, cursor='hand2', font=('Segoe UI',10), padx=10, pady=11)
            b.pack(fill='x', padx=12, pady=2); self.nav_buttons[key] = b

        sep = tk.Frame(self.sidebar, bg=self.theme['border'], height=1); sep.pack(fill='x', padx=22, pady=(20,14))
        tk.Label(self.sidebar, text='QUICK INFO', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI', 8, 'bold')).pack(anchor='w', padx=22)
        tk.Label(self.sidebar, text='Excel + Manual entry', bg=self.theme['sidebar'], fg=self.theme['text'], font=('Segoe UI',9)).pack(anchor='w', padx=22, pady=(8,2))
        tk.Label(self.sidebar, text='Local · private · offline', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI',8)).pack(anchor='w', padx=22)

        bottom = tk.Frame(self.sidebar, bg=self.theme['sidebar']); bottom.pack(side='bottom', fill='x', padx=12, pady=16)
        self.theme_btn = tk.Button(bottom, text='☾  Dark mode', command=self.toggle_theme, relief='flat', bd=0, anchor='w', cursor='hand2', font=('Segoe UI',9), padx=10, pady=9)
        self.theme_btn.pack(fill='x')
        tk.Label(bottom, text='Offline · Local database', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI',8)).pack(anchor='w', padx=10, pady=(7,0))

        self.header = tk.Frame(self.content, bg=self.theme['surface'], height=72, highlightthickness=1, highlightbackground=self.theme['border'])
        self.header.pack(fill='x'); self.header.pack_propagate(False)
        hleft = tk.Frame(self.header, bg=self.theme['surface']); hleft.pack(side='left', padx=28, pady=13)
        self.header_title = tk.Label(hleft, text='', bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI', 16, 'bold')); self.header_title.pack(anchor='w')
        self.header_hint = tk.Label(hleft, text='', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',8)); self.header_hint.pack(anchor='w', pady=(1,0))
        tk.Label(self.header, text='ALI QURESHI  •  نور الحفظ', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',8,'bold')).pack(side='right', padx=28)

        self.page = tk.Frame(self.content, bg=self.theme['bg']); self.page.pack(fill='both', expand=True)
        self.apply_tree_style()

    def style_nav(self, active):
        for key,b in self.nav_buttons.items():
            b.configure(bg=self.theme['accent_soft'] if key == active else self.theme['sidebar'], fg=self.theme['accent'] if key == active else self.theme['text'])
        self.theme_btn.configure(bg=self.theme['sidebar'], fg=self.theme['text'])

    def clear_page(self):
        for w in self.page.winfo_children(): w.destroy()

    def show_page(self, key):
        self.clear_page(); self.style_nav(key)
        titles = {
            'dashboard':('Student Monitoring','Daily overview and local record health.'),
            'import':('Import Daily Report','Select, review and save the formatted Excel report.'),
            'manual':('Manual Entry','Enter a student record directly without Excel.'),
            'report':('Collect Student Report','Find one student and view the complete saved history.'),
            'all':('All Students','Cumulative totals for every student stored locally.'),
        }
        self.header_title.config(text=titles[key][0]); self.header_hint.config(text=titles[key][1])
        if key == 'dashboard': self.build_dashboard()
        elif key == 'import': self.build_import()
        elif key == 'manual': self.build_manual()
        elif key == 'report': self.build_report()
        else: self.build_all()

    def toggle_theme(self):
        self.theme = DARK if self.theme is LIGHT else LIGHT
        for w in self.winfo_children(): w.destroy()
        self.setup_style(); self.build_shell(); self.show_page('dashboard')
        self.theme_btn.config(text='☀  Light mode' if self.theme is DARK else '☾  Dark mode')

    def card(self, parent, **kw):
        return tk.Frame(parent, bg=self.theme['surface'], highlightbackground=self.theme['border'], highlightthickness=1, **kw)

    def section_title(self, parent, title, subtitle=None):
        box = tk.Frame(parent, bg=self.theme['bg']); box.pack(fill='x', pady=(0,15))
        tk.Label(box, text=title, bg=self.theme['bg'], fg=self.theme['text'], font=('Segoe UI', 22, 'bold')).pack(anchor='w')
        if subtitle:
            tk.Label(box, text=subtitle, bg=self.theme['bg'], fg=self.theme['muted'], font=('Segoe UI',10)).pack(anchor='w', pady=(3,0))
        return box

    def stat_card(self, parent, label, value, hint='', accent=None):
        c = self.card(parent); c.pack(side='left', fill='both', expand=True, padx=(0,10))
        if accent:
            tk.Frame(c, bg=accent, width=3).pack(side='left', fill='y')
        inner = tk.Frame(c, bg=self.theme['surface']); inner.pack(fill='both', expand=True, padx=16, pady=14)
        tk.Label(inner, text=value, bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI',20,'bold')).pack(anchor='w')
        tk.Label(inner, text=label, bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',9,'bold')).pack(anchor='w', pady=(2,0))
        if hint:
            tk.Label(inner, text=hint, bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',8)).pack(anchor='w', pady=(5,0))
        return c

    def build_dashboard(self):
        wrap = tk.Frame(self.page, bg=self.theme['bg']); wrap.pack(fill='both', expand=True, padx=30, pady=28)
        self.section_title(wrap, 'Daily overview', f'Today · {date.today().strftime("%d %b %Y")}')
        students, logs, days, today = totals()
        stats = tk.Frame(wrap, bg=self.theme['bg']); stats.pack(fill='x', pady=(0,18))
        self.stat_card(stats, 'Registered students', str(students), 'Students stored locally', self.theme['accent'])
        self.stat_card(stats, 'Absent today', str(today['absent']), 'From today’s imported report', self.theme['danger'])
        self.stat_card(stats, 'Late today', str(today['takheer']), 'Late records today', self.theme['warning'])
        self.stat_card(stats, 'Recorded days', str(days), 'Distinct dates saved', self.theme['success'])

        lower = tk.Frame(wrap, bg=self.theme['bg']); lower.pack(fill='both', expand=True)
        recent = self.card(lower); recent.pack(side='left', fill='both', expand=True, padx=(0,9))
        top = tk.Frame(recent, bg=self.theme['surface']); top.pack(fill='x', padx=20, pady=(18,10))
        tk.Label(top, text='Recent daily records', bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI',12,'bold')).pack(side='left')
        tk.Button(top, text='View all →', command=lambda:self.show_page('all'), bg=self.theme['surface'], fg=self.theme['accent'], relief='flat', bd=0, cursor='hand2', font=('Segoe UI',9,'bold')).pack(side='right')
        cols=('date','students','absent','late','sabaq','sabaqi','manzil')
        tree=ttk.Treeview(recent, columns=cols, show='headings', height=8)
        heads=['Date','Students','Absent','Late','Sabaq','Sabqi','Manzil']
        widths=[110,80,75,70,75,75,75]
        for c,h,w in zip(cols,heads,widths): tree.heading(c,text=h); tree.column(c,width=w,anchor='center')
        for r in recent_days():
            tree.insert('', 'end', values=(r['log_date'],r['students'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']))
        tree.pack(fill='both', expand=True, padx=12, pady=(0,12))

        quick = self.card(lower); quick.pack(side='right', fill='both', expand=False, ipadx=12, padx=(9,0))
        tk.Label(quick, text='Focused workflow', bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI',12,'bold')).pack(anchor='w', padx=22, pady=(20,4))
        tk.Label(quick, text='Everything stays local.', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',9)).pack(anchor='w', padx=22, pady=(0,15))
        for title, desc, command in [
            ('Import report','Select a daily Excel sheet',lambda:self.show_page('import')),
            ('Collect a report','Look up one student',lambda:self.show_page('report')),
            ('All students','Review cumulative totals',lambda:self.show_page('all')),
        ]:
            b=tk.Button(quick, text=title, command=command, bg=self.theme['surface2'], fg=self.theme['text'], activebackground=self.theme['sidebar_hover'], activeforeground=self.theme['text'], relief='flat', bd=0, cursor='hand2', anchor='w', padx=14, pady=10, font=('Segoe UI',9,'bold'))
            b.pack(fill='x', padx=20, pady=4)
            tk.Label(quick,text=desc,bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',8)).pack(anchor='w',padx=34,pady=(0,5))
        tk.Label(quick, text='5 categories · Excel → SQLite → Reports', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',8)).pack(anchor='w', padx=22, pady=(16,20))

    def build_import(self):
        wrap = tk.Frame(self.page, bg=self.theme['bg']); wrap.pack(fill='both', expand=True, padx=30, pady=28)
        self.section_title(wrap, "Import today's report", 'Choose the formatted Excel file. Review detected records before anything is saved.')

        hero = self.card(wrap); hero.pack(fill='x')
        left = tk.Frame(hero, bg=self.theme['surface']); left.pack(side='left', fill='both', expand=True, padx=24, pady=24)
        tk.Label(left, text='Daily Excel report', bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI',13,'bold')).pack(anchor='w')
        tk.Label(left, text='The importer uses the fixed B/C and F/G ranges for all five conditions.\nOther workbook sheets are ignored.', bg=self.theme['surface'], fg=self.theme['muted'], justify='left', font=('Segoe UI',9)).pack(anchor='w', pady=(5,16))
        ttk.Button(left, text='Select Excel File', style='Accent.TButton', command=self.do_import).pack(anchor='w')
        self.status = tk.Label(left, text='No file selected yet.', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',9)); self.status.pack(anchor='w', pady=(13,0))

        side = tk.Frame(hero, bg=self.theme['accent_soft'], width=285); side.pack(side='right', fill='y', padx=1, pady=1); side.pack_propagate(False)
        tk.Label(side, text='IMPORT FLOW', bg=self.theme['accent_soft'], fg=self.theme['accent'], font=('Segoe UI',8,'bold')).pack(anchor='w', padx=20, pady=(22,7))
        for i,t in enumerate(['Select Excel file','Review detected records','Set the report date','Import & Save']):
            tk.Label(side, text=f'{i+1}   {t}', bg=self.theme['accent_soft'], fg=self.theme['text'], font=('Segoe UI',9)).pack(anchor='w', padx=20, pady=5)
        tk.Label(side, text='Future dates are blocked.', bg=self.theme['accent_soft'], fg=self.theme['muted'], font=('Segoe UI',8)).pack(anchor='w', padx=20, pady=(12,0))

        stats = tk.Frame(wrap, bg=self.theme['bg']); stats.pack(fill='x', pady=18)
        self.stat_card(stats, 'Fixed ranges', '5', 'Absent · Late · Sabaq · Sabqi · Manzil', self.theme['accent'])
        self.stat_card(stats, 'Storage', 'SQLite', 'Local database', self.theme['success'])
        self.stat_card(stats, 'Mode', 'Offline', 'No server required', self.theme['warning'])

        info = self.card(wrap); info.pack(fill='x')
        tk.Label(info,text='Template reminder',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',10,'bold')).pack(anchor='w',padx=18,pady=(14,3))
        tk.Label(info,text='Sheet: یومیہ رپورٹ  ·  IDs: B/F  ·  Names: C/G  ·  Date: G1',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=18,pady=(0,14))

    def do_import(self):
        p = filedialog.askopenfilename(title='Select Daily Report', filetypes=[('Excel files','*.xlsx *.xlsm'),('All files','*.*')])
        if not p: return
        try:
            source, recs, counts = read_report(p)
        except Exception as e:
            messagebox.showerror('Could not read report', str(e), parent=self); return
        if not recs:
            messagebox.showwarning('No records found','No student IDs were found in the configured ranges. Make sure IDs are present in B/F.', parent=self); return
        self.status.config(text=f'{Path(p).name} • {len(recs)} students detected', fg=self.theme['accent'])

        w = tk.Toplevel(self); w.title('Noor-ul-Hifz — Review Import'); w.geometry('760x640'); w.minsize(700,600); w.configure(bg=self.theme['bg']); w.transient(self); w.grab_set()
        tk.Frame(w,bg=self.theme['accent'],height=4).pack(fill='x')
        body = tk.Frame(w,bg=self.theme['bg']); body.pack(fill='both',expand=True,padx=28,pady=22)
        tk.Label(body,text='Review Daily Import',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',20,'bold')).pack(anchor='w')
        tk.Label(body,text=Path(p).name,bg=self.theme['bg'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',pady=(3,16))

        date_card = self.card(body); date_card.pack(fill='x',pady=(0,14))
        row = tk.Frame(date_card,bg=self.theme['surface']); row.pack(fill='x',padx=18,pady=16)
        labels = tk.Frame(row,bg=self.theme['surface']); labels.pack(side='left',fill='x',expand=True)
        tk.Label(labels,text='Report date',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',10,'bold')).pack(anchor='w')
        tk.Label(labels,text=f'Excel date: {source}  ·  Defaults to today  ·  Future dates are blocked.',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',pady=(3,0))
        dv = tk.StringVar(value=date.today().isoformat())
        ent = ttk.Entry(row,textvariable=dv,width=18); ent.pack(side='right',ipady=2)

        tk.Label(body,text='Detected records',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',11,'bold')).pack(anchor='w',pady=(4,7))
        grid = tk.Frame(body,bg=self.theme['bg']); grid.pack(fill='x')
        stats=[('Students',len(recs)),('غیر حاضری',counts['absent']),('تاخیر',counts['takheer']),('سبق',counts['sabaq']),('سبقی',counts['sabaqi']),('منزل',counts['manzil'])]
        for i,(lab,val) in enumerate(stats):
            q=self.card(grid); q.grid(row=i//3,column=i%3,sticky='ew',padx=3,pady=3); grid.grid_columnconfigure(i%3,weight=1)
            tk.Label(q,text=str(val),bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',15,'bold')).pack(pady=(10,0))
            tk.Label(q,text=lab,bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(pady=(0,10))
        err=tk.Label(body,text='',bg=self.theme['bg'],fg=self.theme['danger'],font=('Segoe UI',9)); err.pack(anchor='w',pady=(8,0))

        buttons=tk.Frame(body,bg=self.theme['bg']); buttons.pack(side='bottom',fill='x',pady=(18,0))
        ttk.Button(buttons,text='Cancel',command=w.destroy).pack(side='right')
        def save():
            try:
                chosen=datetime.strptime(dv.get().strip(),'%Y-%m-%d').date()
            except ValueError:
                err.config(text='Enter the date in YYYY-MM-DD format.'); return
            if chosen > date.today():
                err.config(text='Future dates are not allowed.'); return
            try:
                save_records(chosen.isoformat(), recs)
                self.status.config(text=f'Saved {len(recs)} students • {chosen.isoformat()}', fg=self.theme['success'])
                w.destroy(); messagebox.showinfo('Import Complete',f'Saved {len(recs)} student records for {chosen.isoformat()}.', parent=self)
            except Exception as e:
                err.config(text='Could not save the report. See the error dialog for details.')
                messagebox.showerror('Save failed', f'{type(e).__name__}: {e}', parent=w)
        ttk.Button(buttons,text='Import & Save',style='Accent.TButton',command=save).pack(side='right',padx=8)

    def build_manual(self):
        wrap=tk.Frame(self.page,bg=self.theme['bg']); wrap.pack(fill='both',expand=True,padx=30,pady=28)
        self.section_title(wrap,'Manual daily entry','Record a student directly. Excel is optional — the same local database is used.')

        top=self.card(wrap); top.pack(fill='x',pady=(0,14))
        inner=tk.Frame(top,bg=self.theme['surface']); inner.pack(fill='x',padx=20,pady=18)
        tk.Label(inner,text='Record date',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9,'bold')).grid(row=0,column=0,sticky='w')
        dv=tk.StringVar(value=date.today().isoformat())
        ttk.Entry(inner,textvariable=dv,width=16).grid(row=1,column=0,sticky='w',pady=(5,0))
        tk.Label(inner,text='Student ID',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9,'bold')).grid(row=0,column=1,sticky='w',padx=(22,0))
        sidv=tk.StringVar(); sid_entry=ttk.Entry(inner,textvariable=sidv,width=22); sid_entry.grid(row=1,column=1,sticky='w',padx=(22,0),pady=(5,0))
        tk.Label(inner,text='Student name',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9,'bold')).grid(row=0,column=2,sticky='w',padx=(22,0))
        namev=tk.StringVar(); ttk.Entry(inner,textvariable=namev,width=32).grid(row=1,column=2,sticky='w',padx=(22,0),pady=(5,0))
        ttk.Button(inner,text='Load',command=lambda:load_student()).grid(row=1,column=3,padx=(12,0),pady=(5,0))
        for i in range(4): inner.grid_columnconfigure(i,weight=0)

        body=self.card(wrap); body.pack(fill='x')
        tk.Label(body,text='Daily status',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',12,'bold')).pack(anchor='w',padx=20,pady=(18,3))
        tk.Label(body,text='Tick every condition that applies. A student can have multiple conditions on the same day.',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=20,pady=(0,14))
        vars={k:tk.IntVar(value=0) for k in LABELS}
        checks=tk.Frame(body,bg=self.theme['surface']); checks.pack(fill='x',padx=18,pady=(0,18))
        for i,(k,label) in enumerate(LABELS.items()):
            box=tk.Frame(checks,bg=self.theme['surface2'],highlightbackground=self.theme['border'],highlightthickness=1)
            box.grid(row=0,column=i,sticky='ew',padx=4); checks.grid_columnconfigure(i,weight=1)
            tk.Checkbutton(box,text=label,variable=vars[k],bg=self.theme['surface2'],fg=self.theme['text'],activebackground=self.theme['surface2'],activeforeground=self.theme['text'],selectcolor=self.theme['surface2'],font=('Segoe UI',9),anchor='w',padx=10,pady=10).pack(fill='x')

        msg=tk.Label(wrap,text='',bg=self.theme['bg'],fg=self.theme['danger'],font=('Segoe UI',9)); msg.pack(anchor='w',pady=(10,0))
        buttons=tk.Frame(wrap,bg=self.theme['bg']); buttons.pack(fill='x',pady=(14,0))
        def clear_form():
            sidv.set(''); namev.set(''); dv.set(date.today().isoformat())
            for v in vars.values(): v.set(0)
            msg.config(text='')
            sid_entry.focus_set()
        def load_student():
            sid=norm(sidv.get())
            if not sid:return
            with conn() as c: row=c.execute('SELECT student_name FROM students WHERE student_id=?',(sid,)).fetchone()
            if row and row['student_name']:
                namev.set(row['student_name']); msg.config(text='Existing student loaded.',fg=self.theme['success'])
            else:
                msg.config(text='New student — enter the name.',fg=self.theme['muted'])
        def save_manual():
            sid=norm(sidv.get()); name=norm(namev.get())
            if not sid:
                msg.config(text='Student ID is required.',fg=self.theme['danger']); sid_entry.focus_set(); return
            if not name:
                msg.config(text='Student name is required.',fg=self.theme['danger']); return
            try: chosen=datetime.strptime(dv.get().strip(),'%Y-%m-%d').date()
            except ValueError:
                msg.config(text='Enter the date as YYYY-MM-DD.',fg=self.theme['danger']); return
            if chosen>date.today():
                msg.config(text='Future dates are not allowed.',fg=self.theme['danger']); return
            rec={'student_id':sid,'student_name':name,**{k:int(v.get()) for k,v in vars.items()}}
            try:
                save_records(chosen.isoformat(),[rec])
                msg.config(text=f'Saved {name} ({sid}) for {chosen.isoformat()}.',fg=self.theme['success'])
                clear_form()
                msg.config(text=f'Saved successfully. {sid} is now stored locally.',fg=self.theme['success'])
                self.status_message='Saved manual entry'
            except Exception as e:
                messagebox.showerror('Save failed',f'{type(e).__name__}: {e}',parent=self)
        ttk.Button(buttons,text='Clear',command=clear_form).pack(side='right')
        ttk.Button(buttons,text='Save Manual Entry',style='Accent.TButton',command=save_manual).pack(side='right',padx=8)
        sid_entry.bind('<Return>',lambda e:load_student())

    def build_report(self):
        f=self.page; wrap=tk.Frame(f,bg=self.theme['bg']); wrap.pack(fill='both',expand=True,padx=30,pady=28)
        self.section_title(wrap, 'Collect Student Report', 'Enter a student ID to view every imported day and cumulative totals.')
        bar=self.card(wrap); bar.pack(fill='x',pady=(0,14)); inner=tk.Frame(bar,bg=self.theme['surface']); inner.pack(fill='x',padx=16,pady=13)
        tk.Label(inner,text='Student ID',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9,'bold')).pack(side='left')
        self.e=ttk.Entry(inner,width=24); self.e.pack(side='left',padx=10); ttk.Button(inner,text='Find',style='Accent.TButton',command=self.find).pack(side='left'); ttk.Button(inner,text='Export Excel',command=self.export).pack(side='left',padx=8)
        self.sum=self.card(wrap); self.sum.pack(fill='x',pady=(0,14)); tk.Label(self.sum,text='Enter an ID and press Find.',bg=self.theme['surface'],fg=self.theme['muted'],anchor='w',justify='left',padx=18,pady=15).pack(fill='x')
        table=self.card(wrap); table.pack(fill='both',expand=True)
        cols=('date','name','absent','late','sabaq','sabaqi','manzil'); self.tree=ttk.Treeview(table,columns=cols,show='headings')
        heads=['Date','Name','غیر حاضری','تاخیر','سبق','سبقی','منزل']
        for x,h in zip(cols,heads): self.tree.heading(x,text=h); self.tree.column(x,width=120 if x=='name' else 95,anchor='center')
        ys=ttk.Scrollbar(table,orient='vertical',command=self.tree.yview); self.tree.configure(yscrollcommand=ys.set); self.tree.pack(side='left',fill='both',expand=True,padx=1,pady=1); ys.pack(side='right',fill='y')
        self.e.bind('<Return>', lambda e:self.find())

    def find(self):
        sid=norm(self.e.get())
        if not sid:return
        with conn() as c:
            st=c.execute('SELECT * FROM students WHERE student_id=?',(sid,)).fetchone()
            rows=c.execute('SELECT * FROM daily_logs WHERE student_id=? ORDER BY log_date',(sid,)).fetchall()
        if not st: messagebox.showinfo('Not Found','No student with this ID is in the local log.',parent=self); return
        self.current=(sid,st['student_name']); [self.tree.delete(x) for x in self.tree.get_children()]; t={k:0 for k in LABELS}
        for r in rows:
            for k in t:t[k]+=r[k]
            self.tree.insert('', 'end', values=(r['log_date'],r['student_name'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']))
        for w in self.sum.winfo_children(): w.destroy()
        tk.Label(self.sum,text=f'ID  {sid}',bg=self.theme['surface'],fg=self.theme['accent'],font=('Segoe UI',9,'bold')).pack(anchor='w',padx=18,pady=(14,0))
        tk.Label(self.sum,text=f'Name  {st["student_name"] or "—"}',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',12,'bold')).pack(anchor='w',padx=18,pady=(2,8))
        tk.Label(self.sum,text=f'غیر حاضری  {t["absent"]}    •    تاخیر  {t["takheer"]}    •    سبق  {t["sabaq"]}    •    سبقی  {t["sabaqi"]}    •    منزل  {t["manzil"]}',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=18,pady=(0,14))

    def export(self):
        if not self.current: messagebox.showinfo('Report','Find a student first.',parent=self); return
        sid,name=self.current; p=filedialog.asksaveasfilename(defaultextension='.xlsx',initialfile=f'student_{sid}_report.xlsx',filetypes=[('Excel files','*.xlsx')])
        if p:
            export_student(sid,name,p); messagebox.showinfo('Exported','Report saved:\n'+p,parent=self)

    def build_all(self):
        f=self.page; wrap=tk.Frame(f,bg=self.theme['bg']); wrap.pack(fill='both',expand=True,padx=30,pady=28)
        head=tk.Frame(wrap,bg=self.theme['bg']); head.pack(fill='x')
        left=tk.Frame(head,bg=self.theme['bg']); left.pack(side='left')
        tk.Label(left,text='All Students',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',22,'bold')).pack(anchor='w')
        tk.Label(left,text='Cumulative totals from all imported daily reports.',bg=self.theme['bg'],fg=self.theme['muted'],font=('Segoe UI',10)).pack(anchor='w',pady=(3,0))
        right=tk.Frame(head,bg=self.theme['bg']); right.pack(side='right',anchor='s')
        ttk.Button(right,text='Refresh',command=self.refresh_all).pack(side='left'); ttk.Button(right,text='Export All Students Excel',style='Accent.TButton',command=self.export_all).pack(side='left',padx=8)
        stats=self.card(wrap); stats.pack(fill='x',pady=18)
        self.all_count=tk.Label(stats,text='0',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',18,'bold')); self.all_count.pack(anchor='w',padx=18,pady=(12,0))
        tk.Label(stats,text='Students stored locally',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=18,pady=(0,12))
        table=self.card(wrap); table.pack(fill='both',expand=True)
        cols=('id','name','absent','late','sabaq','sabaqi','manzil'); self.all_tree=ttk.Treeview(table,columns=cols,show='headings')
        heads={'id':'ID','name':'Name','absent':'غیر حاضری','late':'تاخیر','sabaq':'سبق کا ناغہ','sabaqi':'سبقی کا ناغہ','manzil':'منزل کا ناغہ'}
        widths={'id':110,'name':280,'absent':115,'late':100,'sabaq':135,'sabaqi':135,'manzil':135}
        for c in cols: self.all_tree.heading(c,text=heads[c]); self.all_tree.column(c,width=widths[c],anchor='center')
        ys=ttk.Scrollbar(table,orient='vertical',command=self.all_tree.yview); self.all_tree.configure(yscrollcommand=ys.set); self.all_tree.pack(side='left',fill='both',expand=True,padx=1,pady=1); ys.pack(side='right',fill='y')
        self.refresh_all()

    def all_rows(self):
        with conn() as c:
            return c.execute('''SELECT s.student_id,s.student_name,
                COALESCE(SUM(d.absent),0) absent,COALESCE(SUM(d.takheer),0) takheer,
                COALESCE(SUM(d.sabaq),0) sabaq,COALESCE(SUM(d.sabaqi),0) sabaqi,
                COALESCE(SUM(d.manzil),0) manzil
                FROM students s LEFT JOIN daily_logs d ON d.student_id=s.student_id
                GROUP BY s.student_id,s.student_name ORDER BY s.student_id''').fetchall()

    def refresh_all(self):
        if not hasattr(self,'all_tree'): return
        for x in self.all_tree.get_children(): self.all_tree.delete(x)
        rows=self.all_rows()
        for r in rows:self.all_tree.insert('', 'end', values=(r['student_id'],r['student_name'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']))
        if hasattr(self,'all_count'): self.all_count.config(text=str(len(rows)))

    def export_all(self):
        rows=self.all_rows()
        if not rows: messagebox.showinfo('No Students','There are no students in the local database yet.',parent=self); return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',initialfile='all_students_totals.xlsx',filetypes=[('Excel files','*.xlsx')])
        if not path:return
        wb=Workbook(); ws=wb.active; ws.title='All Students'; ws.append(['ID','Name','غیر حاضری','تاخیر','سبق کا ناغہ','سبقی کا ناغہ','منزل کا ناغہ'])
        for r in rows: ws.append([r['student_id'],r['student_name'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']])
        ws.append([]); ws.append(['TOTAL STUDENTS',len(rows),sum(r['absent'] for r in rows),sum(r['takheer'] for r in rows),sum(r['sabaq'] for r in rows),sum(r['sabaqi'] for r in rows),sum(r['manzil'] for r in rows)])
        ws.freeze_panes='A2'; wb.save(path); messagebox.showinfo('Exported',f'All students report saved:\n{path}',parent=self)


if __name__ == '__main__':
    App().mainloop()
