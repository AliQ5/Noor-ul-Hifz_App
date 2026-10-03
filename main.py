import sqlite3, re
from pathlib import Path
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook, Workbook

APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / 'data' / 'student_logs.db'
RANGES = [
    ('absent', 4, 7),
    ('takheer', 11, 22),
    ('sabaq', 26, 39),
    ('sabaqi', 43, 54),
    ('manzil', 58, 69),
]
LABELS = {
    'absent': 'غیر حاضری',
    'takheer': 'تاخیر',
    'sabaq': 'سبق کا ناغہ',
    'sabaqi': 'سبقی کا ناغہ',
    'manzil': 'منزل کا ناغہ',
}

LIGHT = {
    'bg': '#F5F7FA', 'surface': '#FFFFFF', 'surface2': '#F8FAFC',
    'border': '#E5E7EB', 'text': '#111827', 'muted': '#6B7280',
    'accent': '#2563EB', 'accent_dark': '#1D4ED8', 'accent_soft': '#EFF6FF',
    'success': '#16A34A', 'danger': '#DC2626', 'warning': '#D97706',
    'sidebar': '#FFFFFF', 'sidebar_hover': '#F1F5F9', 'table_head': '#F8FAFC'
}
DARK = {
    'bg': '#0B0F14', 'surface': '#121821', 'surface2': '#171F2A',
    'border': '#263241', 'text': '#F3F4F6', 'muted': '#9CA3AF',
    'accent': '#3B82F6', 'accent_dark': '#2563EB', 'accent_soft': '#172B4D',
    'success': '#22C55E', 'danger': '#F87171', 'warning': '#FBBF24',
    'sidebar': '#0F141B', 'sidebar_hover': '#1A2330', 'table_head': '#171F2A'
}

def norm(v):
    if v is None:
        return ''
    s = str(v).strip()
    return s[:-2] if re.fullmatch(r'-?\d+\.0', s) else s

def conn():
    DB_PATH.parent.mkdir(exist_ok=True)
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c

def init_db():
    with conn() as c:
        c.execute('''CREATE TABLE IF NOT EXISTS daily_logs(
            log_date TEXT, student_id TEXT, student_name TEXT DEFAULT '',
            absent INTEGER DEFAULT 0, takheer INTEGER DEFAULT 0,
            sabaq INTEGER DEFAULT 0, sabaqi INTEGER DEFAULT 0,
            manzil INTEGER DEFAULT 0, created_at TEXT,
            PRIMARY KEY(log_date, student_id))''')
        c.execute('''CREATE TABLE IF NOT EXISTS students(
            student_id TEXT PRIMARY KEY, student_name TEXT DEFAULT '', updated_at TEXT)''')

def read_report(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if 'یومیہ رپورٹ' not in wb.sheetnames:
            raise ValueError("Sheet 'یومیہ رپورٹ' was not found.")
        ws = wb['یومیہ رپورٹ']
        raw = ws['G1'].value
        if isinstance(raw, (datetime, date)):
            source_date = raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat()
        else:
            source_date = norm(raw) or date.today().isoformat()
        found = {}
        counts = {k: 0 for k in LABELS}
        for condition, start, end in RANGES:
            for r in range(start, end + 1):
                for id_col, name_col in [('B', 'C'), ('F', 'G')]:
                    sid = norm(ws[f'{id_col}{r}'].value)
                    name = norm(ws[f'{name_col}{r}'].value)
                    if not sid:
                        continue
                    if sid not in found:
                        found[sid] = {'student_id': sid, 'student_name': name, **{k: 0 for k in LABELS}}
                    elif name and not found[sid]['student_name']:
                        found[sid]['student_name'] = name
                    if not found[sid][condition]:
                        found[sid][condition] = 1
                        counts[condition] += 1
        return source_date, list(found.values()), counts
    finally:
        wb.close()

def save_records(log_date, recs):
    now = datetime.now().isoformat(timespec='seconds')
    with conn() as c:
        for r in recs:
            c.execute('''INSERT INTO daily_logs(
                log_date,student_id,student_name,absent,takheer,sabaq,sabaqi,manzil,created_at)
                VALUES(?,?,?,?,?,?,?,?,?)
                ON CONFLICT(log_date,student_id) DO UPDATE SET
                student_name=CASE WHEN excluded.student_name<>'' THEN excluded.student_name ELSE daily_logs.student_name END,
                absent=MAX(daily_logs.absent,excluded.absent),
                takheer=MAX(daily_logs.takheer,excluded.takheer),
                sabaq=MAX(daily_logs.sabaq,excluded.sabaq),
                sabaqi=MAX(daily_logs.sabaqi,excluded.sabaqi),
                manzil=MAX(daily_logs.manzil,excluded.manzil),
                created_at=excluded.created_at''',
                (log_date, r['student_id'], r['student_name'], r['absent'], r['takheer'], r['sabaq'], r['sabaqi'], r['manzil'], now))
            c.execute('''INSERT INTO students(student_id,student_name,updated_at) VALUES(?,?,?)
                ON CONFLICT(student_id) DO UPDATE SET
                student_name=CASE WHEN excluded.student_name<>'' THEN excluded.student_name ELSE students.student_name END,
                updated_at=excluded.updated_at''',
                (r['student_id'], r['student_name'], now))

def export_student(sid, name, path):
    with conn() as c:
        rows = c.execute('SELECT * FROM daily_logs WHERE student_id=? ORDER BY log_date', (sid,)).fetchall()
    wb = Workbook(); ws = wb.active; ws.title = 'Student Report'
    ws.append(['ID','Name','Date','غیر حاضری','تاخیر','سبق کا ناغہ','سبقی کا ناغہ','منزل کا ناغہ'])
    for r in rows:
        ws.append([r['student_id'], r['student_name'] or name, r['log_date'], r['absent'], r['takheer'], r['sabaq'], r['sabaqi'], r['manzil']])
    ws.append([])
    ws.append(['Totals','','',sum(r['absent'] for r in rows),sum(r['takheer'] for r in rows),sum(r['sabaq'] for r in rows),sum(r['sabaqi'] for r in rows),sum(r['manzil'] for r in rows)])
    ws.freeze_panes = 'A2'; wb.save(path)

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Student Daily Log')
        self.geometry('1240x760'); self.minsize(1050,680)
        self.theme = LIGHT
        self.current = None
        self.setup_style(); init_db(); self.build_shell(); self.show_page('import')

    def setup_style(self):
        st = ttk.Style(self)
        try: st.theme_use('clam')
        except tk.TclError: pass
        st.configure('TButton', font=('Segoe UI', 10), padding=(13, 8))
        st.configure('Accent.TButton', font=('Segoe UI', 10, 'bold'), padding=(15, 9))
        st.configure('TEntry', font=('Segoe UI', 10), padding=8)
        st.configure('Treeview', font=('Segoe UI', 10), rowheight=40)
        st.configure('Treeview.Heading', font=('Segoe UI', 10, 'bold'), padding=9)

    def apply_tree_style(self):
        st = ttk.Style(self)
        st.configure('Treeview', background=self.theme['surface'], fieldbackground=self.theme['surface'], foreground=self.theme['text'], rowheight=42, borderwidth=0)
        st.map('Treeview', background=[('selected', self.theme['accent'])], foreground=[('selected', '#FFFFFF')])
        st.configure('Treeview.Heading', background=self.theme['table_head'], foreground=self.theme['muted'], font=('Segoe UI', 10, 'bold'))
        st.configure('TButton', background=self.theme['surface'], foreground=self.theme['text'])
        st.map('TButton', background=[('active', self.theme['sidebar_hover'])])
        st.configure('Accent.TButton', background=self.theme['accent'], foreground='#FFFFFF')
        st.map('Accent.TButton', background=[('active', self.theme['accent_dark'])])
        st.configure('TEntry', fieldbackground=self.theme['surface'], foreground=self.theme['text'])

    def build_shell(self):
        self.configure(bg=self.theme['bg'])
        self.sidebar = tk.Frame(self, bg=self.theme['sidebar'], width=225, highlightthickness=1, highlightbackground=self.theme['border'])
        self.sidebar.pack(side='left', fill='y'); self.sidebar.pack_propagate(False)
        self.content = tk.Frame(self, bg=self.theme['bg']); self.content.pack(side='left', fill='both', expand=True)

        brand = tk.Frame(self.sidebar, bg=self.theme['sidebar']); brand.pack(fill='x', padx=18, pady=(22,28))
        logo = tk.Frame(brand, bg=self.theme['accent'], width=38, height=38); logo.pack(side='left'); logo.pack_propagate(False)
        tk.Label(logo, text='S', bg=self.theme['accent'], fg='white', font=('Segoe UI', 17, 'bold')).pack(expand=True)
        tk.Label(brand, text='Student Daily\nLog', bg=self.theme['sidebar'], fg=self.theme['text'], justify='left', font=('Segoe UI', 12, 'bold')).pack(side='left', padx=10)

        tk.Label(self.sidebar, text='WORKSPACE', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI', 8, 'bold')).pack(anchor='w', padx=22, pady=(0,8))
        self.nav_buttons = {}
        for key, title, icon in [('import','Import Daily Report','↓'),('report','Collect Report','▤'),('all','All Students','≡')]:
            b = tk.Button(self.sidebar, text=f'  {icon}   {title}', command=lambda k=key:self.show_page(k), anchor='w', relief='flat', bd=0, cursor='hand2', font=('Segoe UI',10), padx=10, pady=10)
            b.pack(fill='x', padx=12, pady=2); self.nav_buttons[key] = b

        bottom = tk.Frame(self.sidebar, bg=self.theme['sidebar']); bottom.pack(side='bottom', fill='x', padx=12, pady=16)
        self.theme_btn = tk.Button(bottom, text='☾  Dark mode', command=self.toggle_theme, relief='flat', bd=0, anchor='w', cursor='hand2', font=('Segoe UI',9), padx=10, pady=9)
        self.theme_btn.pack(fill='x')
        tk.Label(bottom, text='Offline • Local database', bg=self.theme['sidebar'], fg=self.theme['muted'], font=('Segoe UI',8)).pack(anchor='w', padx=10, pady=(7,0))

        self.header = tk.Frame(self.content, bg=self.theme['surface'], height=66, highlightthickness=1, highlightbackground=self.theme['border'])
        self.header.pack(fill='x'); self.header.pack_propagate(False)
        self.header_title = tk.Label(self.header, text='', bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI', 16, 'bold'))
        self.header_title.pack(side='left', padx=26)
        self.header_hint = tk.Label(self.header, text='', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',9))
        self.header_hint.pack(side='left', padx=8)
        self.page = tk.Frame(self.content, bg=self.theme['bg']); self.page.pack(fill='both', expand=True)
        self.apply_tree_style()

    def style_nav(self, active):
        for key,b in self.nav_buttons.items():
            b.configure(bg=self.theme['accent_soft'] if key == active else self.theme['sidebar'], fg=self.theme['accent'] if key == active else self.theme['text'])
        self.theme_btn.configure(bg=self.theme['sidebar'], fg=self.theme['text'])

    def clear_page(self):
        for w in self.page.winfo_children(): w.destroy()

    def show_page(self, key):
        self.clear_page(); self.style_nav(key)
        titles = {'import':('Import Daily Report','Review and save a formatted daily Excel report.'), 'report':('Collect Student Report','View a student\'s complete attendance and performance history.'), 'all':('All Students','Cumulative totals for every student stored locally.')}
        self.header_title.config(text=titles[key][0]); self.header_hint.config(text=titles[key][1])
        if key == 'import': self.build_import()
        elif key == 'report': self.build_report()
        else: self.build_all()

    def toggle_theme(self):
        self.theme = DARK if self.theme is LIGHT else LIGHT
        # Rebuild the shell for a consistent full-theme refresh while preserving the database.
        for w in self.winfo_children(): w.destroy()
        self.setup_style(); self.build_shell(); self.show_page('import')
        self.theme_btn.config(text='☀  Light mode' if self.theme is DARK else '☾  Dark mode')

    def card(self, parent, **kw):
        return tk.Frame(parent, bg=self.theme['surface'], highlightbackground=self.theme['border'], highlightthickness=1, **kw)

    def build_import(self):
        f = self.page
        wrap = tk.Frame(f, bg=self.theme['bg']); wrap.pack(fill='both', expand=True, padx=30, pady=28)
        tk.Label(wrap, text='Import today\'s report', bg=self.theme['bg'], fg=self.theme['text'], font=('Segoe UI', 23, 'bold')).pack(anchor='w')
        tk.Label(wrap, text='Choose the formatted Excel file. You will review the detected records before anything is saved.', bg=self.theme['bg'], fg=self.theme['muted'], font=('Segoe UI',10)).pack(anchor='w', pady=(4,20))

        hero = self.card(wrap); hero.pack(fill='x')
        left = tk.Frame(hero, bg=self.theme['surface']); left.pack(side='left', fill='both', expand=True, padx=24, pady=24)
        tk.Label(left, text='Daily Excel report', bg=self.theme['surface'], fg=self.theme['text'], font=('Segoe UI',13,'bold')).pack(anchor='w')
        tk.Label(left, text='The importer uses the fixed B/C and F/G ranges for all five conditions.\nOther workbook sheets are ignored.', bg=self.theme['surface'], fg=self.theme['muted'], justify='left', font=('Segoe UI',9)).pack(anchor='w', pady=(5,16))
        ttk.Button(left, text='Select Excel File', style='Accent.TButton', command=self.do_import).pack(anchor='w')
        self.status = tk.Label(left, text='No file selected yet.', bg=self.theme['surface'], fg=self.theme['muted'], font=('Segoe UI',9))
        self.status.pack(anchor='w', pady=(13,0))
        side = tk.Frame(hero, bg=self.theme['accent_soft'], width=260); side.pack(side='right', fill='y', padx=1, pady=1); side.pack_propagate(False)
        tk.Label(side, text='IMPORT FLOW', bg=self.theme['accent_soft'], fg=self.theme['accent'], font=('Segoe UI',8,'bold')).pack(anchor='w', padx=20, pady=(22,7))
        for i,t in enumerate(['Select Excel file','Review detected records','Set the report date','Import & Save']):
            tk.Label(side, text=f'{i+1}   {t}', bg=self.theme['accent_soft'], fg=self.theme['text'], font=('Segoe UI',9)).pack(anchor='w', padx=20, pady=5)

        stats = tk.Frame(wrap, bg=self.theme['bg']); stats.pack(fill='x', pady=18)
        for title, value in [('Fixed ranges','5'),('Database','SQLite'),('Mode','Offline')]:
            c = self.card(stats); c.pack(side='left', fill='x', expand=True, padx=(0,10))
            tk.Label(c,text=value,bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',15,'bold')).pack(anchor='w',padx=16,pady=(12,0))
            tk.Label(c,text=title,bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=16,pady=(0,12))

    def do_import(self):
        p = filedialog.askopenfilename(title='Select Daily Report', filetypes=[('Excel files','*.xlsx *.xlsm'),('All files','*.*')])
        if not p: return
        try:
            source, recs, counts = read_report(p)
        except Exception as e:
            messagebox.showerror('Could not read report', str(e), parent=self); return
        if not recs:
            messagebox.showwarning('No records found','No student IDs were found in the configured ranges. Make sure IDs are present in B/F.', parent=self); return

        w = tk.Toplevel(self); w.title('Review Daily Import'); w.geometry('700x600'); w.minsize(650,560); w.configure(bg=self.theme['bg']); w.transient(self); w.grab_set()
        top = tk.Frame(w,bg=self.theme['accent'],height=5); top.pack(fill='x')
        body = tk.Frame(w,bg=self.theme['bg']); body.pack(fill='both',expand=True,padx=24,pady=20)
        tk.Label(body,text='Review Daily Import',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',20,'bold')).pack(anchor='w')
        tk.Label(body,text=Path(p).name,bg=self.theme['bg'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',pady=(3,16))

        date_card = self.card(body); date_card.pack(fill='x',pady=(0,12))
        row = tk.Frame(date_card,bg=self.theme['surface']); row.pack(fill='x',padx=18,pady=16)
        labels = tk.Frame(row,bg=self.theme['surface']); labels.pack(side='left',fill='x',expand=True)
        tk.Label(labels,text='Report date',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',10,'bold')).pack(anchor='w')
        tk.Label(labels,text='Defaults to today. Future dates are not allowed.',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',pady=(3,0))
        dv = tk.StringVar(value=date.today().isoformat())
        ent = ttk.Entry(row,textvariable=dv,width=18); ent.pack(side='right',ipady=2)
        ent.focus_set()

        tk.Label(body,text='Detected records',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',11,'bold')).pack(anchor='w',pady=(5,7))
        grid = tk.Frame(body,bg=self.theme['bg']); grid.pack(fill='x')
        stats=[('Students',len(recs)),('غیر حاضری',counts['absent']),('تاخیر',counts['takheer']),('سبق',counts['sabaq']),('سبقی',counts['sabaqi']),('منزل',counts['manzil'])]
        for i,(lab,val) in enumerate(stats):
            q=self.card(grid); q.grid(row=i//3,column=i%3,sticky='ew',padx=3,pady=3); grid.grid_columnconfigure(i%3,weight=1)
            tk.Label(q,text=str(val),bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',15,'bold')).pack(pady=(10,0))
            tk.Label(q,text=lab,bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(pady=(0,10))
        err=tk.Label(body,text='',bg=self.theme['bg'],fg=self.theme['danger'],font=('Segoe UI',9)); err.pack(anchor='w',pady=(8,0))

        buttons=tk.Frame(body,bg=self.theme['bg']); buttons.pack(side='bottom',fill='x',pady=(18,0))
        ttk.Button(buttons,text='Cancel',command=w.destroy).pack(side='right')
        def save():
            try:
                chosen=datetime.strptime(dv.get().strip(),'%Y-%m-%d').date()
            except ValueError:
                err.config(text='Enter the date in YYYY-MM-DD format.'); return
            if chosen > date.today():
                err.config(text='Future dates are not allowed.'); return
            try:
                save_records(chosen.isoformat(), recs)
                self.status.config(text=f'Saved {len(recs)} students • {chosen.isoformat()}', fg=self.theme['success'])
                w.destroy(); messagebox.showinfo('Import Complete',f'Saved {len(recs)} student records for {chosen.isoformat()}.', parent=self)
                if hasattr(self,'all_tree'): self.refresh_all()
            except Exception as e:
                err.config(text='Could not save the report. See the error dialog for details.')
                messagebox.showerror('Save failed', f'{type(e).__name__}: {e}', parent=w)
        ttk.Button(buttons,text='Import & Save',style='Accent.TButton',command=save).pack(side='right',padx=8)

    def build_report(self):
        f=self.page; wrap=tk.Frame(f,bg=self.theme['bg']); wrap.pack(fill='both',expand=True,padx=30,pady=28)
        tk.Label(wrap,text='Collect Student Report',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',22,'bold')).pack(anchor='w')
        tk.Label(wrap,text='Enter a student ID to view every imported day and cumulative totals.',bg=self.theme['bg'],fg=self.theme['muted'],font=('Segoe UI',10)).pack(anchor='w',pady=(3,16))
        bar=self.card(wrap); bar.pack(fill='x',pady=(0,14)); inner=tk.Frame(bar,bg=self.theme['surface']); inner.pack(fill='x',padx=16,pady=13)
        tk.Label(inner,text='Student ID',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9,'bold')).pack(side='left')
        self.e=ttk.Entry(inner,width=22); self.e.pack(side='left',padx=10); ttk.Button(inner,text='Find',style='Accent.TButton',command=self.find).pack(side='left'); ttk.Button(inner,text='Export Excel',command=self.export).pack(side='left',padx=8)
        self.sum=self.card(wrap); self.sum.pack(fill='x',pady=(0,14)); tk.Label(self.sum,text='Enter an ID and press Find.',bg=self.theme['surface'],fg=self.theme['muted'],anchor='w',justify='left',padx=18,pady=15).pack(fill='x')
        table=self.card(wrap); table.pack(fill='both',expand=True)
        cols=('date','name','absent','late','sabaq','sabaqi','manzil'); self.tree=ttk.Treeview(table,columns=cols,show='headings')
        heads=['Date','Name','غیر حاضری','تاخیر','سبق','سبقی','منزل']
        for x,h in zip(cols,heads): self.tree.heading(x,text=h); self.tree.column(x,width=120 if x=='name' else 95,anchor='center')
        ys=ttk.Scrollbar(table,orient='vertical',command=self.tree.yview); self.tree.configure(yscrollcommand=ys.set); self.tree.pack(side='left',fill='both',expand=True,padx=1,pady=1); ys.pack(side='right',fill='y')

    def find(self):
        sid=norm(self.e.get())
        if not sid:return
        with conn() as c:
            st=c.execute('SELECT * FROM students WHERE student_id=?',(sid,)).fetchone()
            rows=c.execute('SELECT * FROM daily_logs WHERE student_id=? ORDER BY log_date',(sid,)).fetchall()
        if not st: messagebox.showinfo('Not Found','No student with this ID is in the local log.',parent=self); return
        self.current=(sid,st['student_name']); [self.tree.delete(x) for x in self.tree.get_children()]; t={k:0 for k in LABELS}
        for r in rows:
            for k in t:t[k]+=r[k]
            self.tree.insert('', 'end', values=(r['log_date'],r['student_name'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']))
        for w in self.sum.winfo_children(): w.destroy()
        tk.Label(self.sum,text=f'ID  {sid}',bg=self.theme['surface'],fg=self.theme['accent'],font=('Segoe UI',10,'bold')).pack(anchor='w',padx=18,pady=(14,0))
        tk.Label(self.sum,text=f'Name  {st["student_name"] or "—"}',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',12,'bold')).pack(anchor='w',padx=18,pady=(2,8))
        tk.Label(self.sum,text=f'غیر حاضری  {t["absent"]}    •    تاخیر  {t["takheer"]}    •    سبق  {t["sabaq"]}    •    سبقی  {t["sabaqi"]}    •    منزل  {t["manzil"]}',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=18,pady=(0,14))

    def export(self):
        if not self.current: messagebox.showinfo('Report','Find a student first.',parent=self); return
        sid,name=self.current; p=filedialog.asksaveasfilename(defaultextension='.xlsx',initialfile=f'student_{sid}_report.xlsx',filetypes=[('Excel files','*.xlsx')])
        if p: export_student(sid,name,p); messagebox.showinfo('Exported','Report saved:\n'+p,parent=self)

    def build_all(self):
        f=self.page; wrap=tk.Frame(f,bg=self.theme['bg']); wrap.pack(fill='both',expand=True,padx=30,pady=28)
        head=tk.Frame(wrap,bg=self.theme['bg']); head.pack(fill='x')
        left=tk.Frame(head,bg=self.theme['bg']); left.pack(side='left')
        tk.Label(left,text='All Students',bg=self.theme['bg'],fg=self.theme['text'],font=('Segoe UI',22,'bold')).pack(anchor='w')
        tk.Label(left,text='Cumulative totals from all imported daily reports.',bg=self.theme['bg'],fg=self.theme['muted'],font=('Segoe UI',10)).pack(anchor='w',pady=(3,0))
        right=tk.Frame(head,bg=self.theme['bg']); right.pack(side='right',anchor='s')
        ttk.Button(right,text='Refresh',command=self.refresh_all).pack(side='left'); ttk.Button(right,text='Export All Students Excel',style='Accent.TButton',command=self.export_all).pack(side='left',padx=8)
        stats=self.card(wrap); stats.pack(fill='x',pady=18)
        self.all_count=tk.Label(stats,text='0',bg=self.theme['surface'],fg=self.theme['text'],font=('Segoe UI',18,'bold')); self.all_count.pack(anchor='w',padx=18,pady=(12,0))
        tk.Label(stats,text='Students stored locally',bg=self.theme['surface'],fg=self.theme['muted'],font=('Segoe UI',9)).pack(anchor='w',padx=18,pady=(0,12))
        table=self.card(wrap); table.pack(fill='both',expand=True)
        cols=('id','name','absent','late','sabaq','sabaqi','manzil'); self.all_tree=ttk.Treeview(table,columns=cols,show='headings')
        heads={'id':'ID','name':'Name','absent':'غیر حاضری','late':'تاخیر','sabaq':'سبق کا ناغہ','sabaqi':'سبقی کا ناغہ','manzil':'منزل کا ناغہ'}
        widths={'id':110,'name':280,'absent':115,'late':100,'sabaq':135,'sabaqi':135,'manzil':135}
        for c in cols: self.all_tree.heading(c,text=heads[c]); self.all_tree.column(c,width=widths[c],anchor='center')
        ys=ttk.Scrollbar(table,orient='vertical',command=self.all_tree.yview); self.all_tree.configure(yscrollcommand=ys.set); self.all_tree.pack(side='left',fill='both',expand=True,padx=1,pady=1); ys.pack(side='right',fill='y')
        self.refresh_all()

    def all_rows(self):
        with conn() as c:
            return c.execute('''SELECT s.student_id,s.student_name,
                COALESCE(SUM(d.absent),0) absent,COALESCE(SUM(d.takheer),0) takheer,
                COALESCE(SUM(d.sabaq),0) sabaq,COALESCE(SUM(d.sabaqi),0) sabaqi,
                COALESCE(SUM(d.manzil),0) manzil
                FROM students s LEFT JOIN daily_logs d ON d.student_id=s.student_id
                GROUP BY s.student_id,s.student_name ORDER BY s.student_id''').fetchall()

    def refresh_all(self):
        if not hasattr(self,'all_tree'): return
        for x in self.all_tree.get_children(): self.all_tree.delete(x)
        rows=self.all_rows()
        for r in rows:self.all_tree.insert('', 'end', values=(r['student_id'],r['student_name'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']))
        if hasattr(self,'all_count'): self.all_count.config(text=str(len(rows)))

    def export_all(self):
        rows=self.all_rows()
        if not rows: messagebox.showinfo('No Students','There are no students in the local database yet.',parent=self); return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',initialfile='all_students_totals.xlsx',filetypes=[('Excel files','*.xlsx')])
        if not path:return
        wb=Workbook(); ws=wb.active; ws.title='All Students'; ws.append(['ID','Name','غیر حاضری','تاخیر','سبق کا ناغہ','سبقی کا ناغہ','منزل کا ناغہ'])
        for r in rows: ws.append([r['student_id'],r['student_name'],r['absent'],r['takheer'],r['sabaq'],r['sabaqi'],r['manzil']])
        ws.append([]); ws.append(['TOTAL STUDENTS',len(rows),sum(r['absent'] for r in rows),sum(r['takheer'] for r in rows),sum(r['sabaq'] for r in rows),sum(r['sabaqi'] for r in rows),sum(r['manzil'] for r in rows)])
        ws.freeze_panes='A2'; wb.save(path); messagebox.showinfo('Exported',f'All students report saved:\n{path}',parent=self)

if __name__ == '__main__':
    App().mainloop()
